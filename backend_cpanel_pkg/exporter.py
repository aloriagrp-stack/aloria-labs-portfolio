"""
Aloria Hunter - Outreach Data Exporter Module
Exports all sent emails and complete lead intelligence to:
1. Agent JSON (sent_emails_for_agent.json)
2. Agent CSV (sent_emails_for_agent.csv)
3. Plain Text Unique Email List (sent_emails_unique_list.txt)
4. Full Excel Workbook (outreach_campaign_sent_emails.xlsx)
5. Standard CSV (outreach_campaign_sent_emails.csv)
"""

import sys
import os

# Force UTF-8 stream handling with replacement to prevent Windows cp1252 charmap crashes
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")  # type: ignore
    except Exception:
        pass
if hasattr(sys.stderr, "reconfigure"):
    try:
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")  # type: ignore
    except Exception:
        pass

import csv
import json
import argparse
import urllib.parse
from datetime import datetime
from pathlib import Path

from typing import TYPE_CHECKING
if TYPE_CHECKING:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

# Optional openpyxl import with graceful fallback
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# Ensure hunter directory is in sys.path
BASE_DIR = Path(__file__).resolve().parent
PARENT_DIR = BASE_DIR.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

import db

def clean_email(em):
    if not em:
        return ""
    em = urllib.parse.unquote(str(em)).strip().lower()
    em = em.strip(" -_'\"<>:,;/\\")
    return em

def safe_write_text_or_csv(target_path: Path, write_fn, mode="w", encoding="utf-8", newline=None):
    """
    Safely writes to target_path. If file is locked (e.g. open in Excel),
    falls back to a timestamped file name instead of crashing.
    """
    try:
        with open(target_path, mode, encoding=encoding, newline=newline) as f:
            write_fn(f)
        return target_path
    except PermissionError:
        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        alt_path = target_path.parent / f"{target_path.stem}_{ts}{target_path.suffix}"
        try:
            with open(alt_path, mode, encoding=encoding, newline=newline) as f:
                write_fn(f)
            print(f"[!] Warning: '{target_path.name}' is currently locked. Saved to '{alt_path.name}' instead.")
            return alt_path
        except Exception as e:
            print(f"[-] Error writing fallback file '{alt_path.name}': {e}")
            return target_path

def generate_outreach_exports(output_dir: Path | str | None = None, business_id: str | None = None) -> dict:
    """
    Generates all outreach export files (JSON, CSV, TXT, Excel).
    - output_dir: Destination directory (defaults to PARENT_DIR / workspace root)
    - business_id: Optional filter for a specific business/campaign
    """
    out_dir = Path(output_dir) if output_dir else PARENT_DIR
    out_dir.mkdir(parents=True, exist_ok=True)

    conn = db.get_connection()
    try:
        c = conn.cursor()

        sql = """
        SELECT 
            o.id AS log_id,
            o.lead_id,
            COALESCE(o.business_name, l.business_name, '') AS business_name,
            LOWER(TRIM(COALESCE(o.recipient_email, ''))) AS recipient_email,
            COALESCE(o.sender_email, '') AS sender_email,
            COALESCE(NULLIF(o.business_id, ''), NULLIF(l.business_id, ''), 'aloria_labs') AS campaign,
            COALESCE(o.pitch_type, 'INITIAL') AS pitch_type,
            COALESCE(o.subject, '') AS subject,
            COALESCE(o.status, 'SENT') AS delivery_status,
            COALESCE(o.sent_at, '') AS sent_at,
            COALESCE(l.city, '') AS city,
            COALESCE(l.country, '') AS country,
            COALESCE(l.niche, '') AS niche,
            COALESCE(l.website_url, '') AS website_url,
            COALESCE(l.phone, '') AS phone,
            COALESCE(l.rating, '') AS rating,
            COALESCE(l.reviews_count, '') AS reviews_count,
            COALESCE(l.follow_up_count, 0) AS follow_up_count,
            COALESCE(l.last_follow_up_at, '') AS last_follow_up_at,
            COALESCE(l.initial_email_sent_at, '') AS initial_email_sent_at,
            COALESCE(l.status, 'SENT') AS current_lead_status,
            COALESCE(l.audit_summary, '') AS audit_summary
        FROM outreach_logs o
        LEFT JOIN leads l ON o.lead_id = l.id
        WHERE UPPER(TRIM(o.status)) = 'SENT'
        """
        params = []
        if business_id:
            sql += " AND (o.business_id = ? OR l.business_id = ?)"
            params.extend([business_id, business_id])

        sql += " ORDER BY o.id ASC"
        c.execute(sql, params)
        rows = [dict(r) for r in c.fetchall()]

        # 2. Leads table contacted backup (catches any leads where emails were sent)
        leads_sql = """
        SELECT 
            id AS lead_id,
            COALESCE(business_name, '') AS business_name,
            LOWER(TRIM(email)) AS recipient_email,
            COALESCE(business_id, 'aloria_labs') AS campaign,
            COALESCE(status, 'SENT') AS current_lead_status,
            COALESCE(initial_email_sent_at, '') AS initial_email_sent_at,
            COALESCE(last_follow_up_at, '') AS last_follow_up_at,
            COALESCE(follow_up_count, 0) AS follow_up_count,
            COALESCE(city, '') AS city,
            COALESCE(country, '') AS country,
            COALESCE(niche, '') AS niche,
            COALESCE(website_url, '') AS website_url,
            COALESCE(phone, '') AS phone,
            COALESCE(rating, '') AS rating,
            COALESCE(reviews_count, '') AS reviews_count,
            COALESCE(audit_summary, '') AS audit_summary
        FROM leads
        WHERE initial_email_sent_at IS NOT NULL 
          AND TRIM(initial_email_sent_at) != ''
          AND email IS NOT NULL AND TRIM(email) != ''
          AND LOWER(TRIM(email)) LIKE '%@%.%'
        """
        leads_params = []
        if business_id:
            leads_sql += " AND business_id = ?"
            leads_params.append(business_id)
        leads_sql += " ORDER BY id ASC"

        c.execute(leads_sql, leads_params)
        leads_contacted = [dict(r) for r in c.fetchall()]
    finally:
        try:
            conn.close()
        except Exception:
            pass

    # Clean emails
    for r in rows:
        r["recipient_email_clean"] = clean_email(r["recipient_email"])
    for r in leads_contacted:
        r["recipient_email_clean"] = clean_email(r["recipient_email"])

    # Unique valid emails set
    unique_emails_set = set()
    for r in rows:
        em = r["recipient_email_clean"]
        if em and "@" in em and "." in em:
            unique_emails_set.add(em)
    for r in leads_contacted:
        em = r["recipient_email_clean"]
        if em and "@" in em and "." in em:
            unique_emails_set.add(em)

    sorted_unique_emails = sorted(list(unique_emails_set))

    # Fast O(1) indices to avoid multi-million iteration scans
    first_log_by_email = {}
    email_history_map = {}
    for r in rows:
        em = r["recipient_email_clean"]
        if not em:
            continue
        if em not in first_log_by_email:
            first_log_by_email[em] = r
        if em not in email_history_map:
            email_history_map[em] = []
        email_history_map[em].append({
            "log_id": r["log_id"],
            "pitch_type": r["pitch_type"],
            "subject": r["subject"],
            "sender_email": r["sender_email"],
            "campaign": r["campaign"],
            "sent_at": r["sent_at"]
        })

    leads_by_email = {}
    for r in leads_contacted:
        em = r["recipient_email_clean"]
        if em and em not in leads_by_email:
            leads_by_email[em] = r

    contacted_leads_structured = []
    for em in sorted_unique_emails:
        ld = leads_by_email.get(em, {})
        log_sample = first_log_by_email.get(em, {})
        history = email_history_map.get(em, [])
        first_event = history[0] if history else {}
        last_event = history[-1] if history else {}

        contacted_leads_structured.append({
            "recipient_email": em,
            "business_name": ld.get("business_name") or log_sample.get("business_name", ""),
            "campaign": ld.get("campaign") or log_sample.get("campaign", ""),
            "city": ld.get("city") or log_sample.get("city", ""),
            "country": ld.get("country") or log_sample.get("country", ""),
            "niche": ld.get("niche") or log_sample.get("niche", ""),
            "website_url": ld.get("website_url") or log_sample.get("website_url", ""),
            "phone": ld.get("phone") or log_sample.get("phone", ""),
            "rating": ld.get("rating") or log_sample.get("rating", ""),
            "reviews_count": ld.get("reviews_count") or log_sample.get("reviews_count", ""),
            "status": ld.get("current_lead_status") or log_sample.get("current_lead_status", "SENT"),
            "initial_email_sent_at": ld.get("initial_email_sent_at") or first_event.get("sent_at", ""),
            "follow_up_count": ld.get("follow_up_count") if ld.get("follow_up_count") is not None else (len(history) - 1 if len(history) > 1 else 0),
            "last_follow_up_at": ld.get("last_follow_up_at") or (last_event.get("sent_at", "") if len(history) > 1 else ""),
            "audit_summary": ld.get("audit_summary") or log_sample.get("audit_summary", ""),
            "total_emails_sent": len(history),
            "dispatch_history": history
        })

    # Destination paths
    xlsx_path = out_dir / "outreach_campaign_sent_emails.xlsx"
    csv_path = out_dir / "outreach_campaign_sent_emails.csv"
    agent_json_path = out_dir / "sent_emails_for_agent.json"
    agent_csv_path = out_dir / "sent_emails_for_agent.csv"
    agent_txt_path = out_dir / "sent_emails_unique_list.txt"

    # 1. Agent JSON Export
    pitch_counts = {}
    campaign_counts = {}
    for r in rows:
        pt = r.get("pitch_type") or "UNKNOWN"
        pitch_counts[pt] = pitch_counts.get(pt, 0) + 1
        cp = r.get("campaign") or "UNKNOWN"
        campaign_counts[cp] = campaign_counts.get(cp, 0) + 1

    summary = {
        "total_sent_email_events": len(rows),
        "total_unique_contacted_emails": len(sorted_unique_emails),
        "campaigns_breakdown": campaign_counts,
        "pitch_types_breakdown": pitch_counts,
        "generated_at": datetime.now().isoformat()
    }

    def write_json(f):
        json.dump({
            "metadata": summary,
            "unique_emails_list": sorted_unique_emails,
            "leads_data": contacted_leads_structured,
            "raw_sent_logs": rows
        }, f, indent=2, ensure_ascii=False)

    agent_json_path = safe_write_text_or_csv(agent_json_path, write_json, mode="w", encoding="utf-8")

    # 2. Agent CSV Export
    agent_csv_headers = [
        "Recipient Email", "Business Name", "Campaign", "Pitch Type", "Subject",
        "Sender Email", "Sent At", "Delivery Status", "City", "Country", "Niche",
        "Website URL", "Phone", "Rating", "Reviews Count", "Follow-up Count",
        "Lead Status", "Audit Summary"
    ]
    def write_agent_csv(f):
        writer = csv.writer(f)
        writer.writerow(agent_csv_headers)
        for r in rows:
            writer.writerow([
                r.get("recipient_email_clean", ""), r.get("business_name", ""), r.get("campaign", ""), r.get("pitch_type", ""),
                r.get("subject", ""), r.get("sender_email", ""), r.get("sent_at", ""), r.get("delivery_status", ""),
                r.get("city", ""), r.get("country", ""), r.get("niche", ""), r.get("website_url", ""), r.get("phone", ""),
                r.get("rating", ""), r.get("reviews_count", ""), r.get("follow_up_count", 0), r.get("current_lead_status", ""),
                r.get("audit_summary", "")
            ])

    agent_csv_path = safe_write_text_or_csv(agent_csv_path, write_agent_csv, mode="w", encoding="utf-8-sig", newline="")

    # 3. Agent Plain Text Email List
    def write_txt(f):
        for em in sorted_unique_emails:
            f.write(f"{em}\n")

    agent_txt_path = safe_write_text_or_csv(agent_txt_path, write_txt, mode="w", encoding="utf-8")

    # 4. Standard CSV
    headers = [
        "Log ID", "Lead ID", "Business Name", "Recipient Email", "Sender Email",
        "Campaign", "Pitch Type", "Subject", "Delivery Status", "Sent At",
        "City", "Country", "Niche", "Website URL", "Phone", "Rating",
        "Reviews Count", "Follow-up Count", "Last Follow-up At", "Current Lead Status"
    ]
    def write_standard_csv(f):
        writer = csv.writer(f)
        writer.writerow(headers)
        for r in rows:
            writer.writerow([
                r.get("log_id", ""), r.get("lead_id", ""), r.get("business_name", ""), r.get("recipient_email_clean", ""), r.get("sender_email", ""),
                r.get("campaign", ""), r.get("pitch_type", ""), r.get("subject", ""), r.get("delivery_status", ""), r.get("sent_at", ""),
                r.get("city", ""), r.get("country", ""), r.get("niche", ""), r.get("website_url", ""), r.get("phone", ""), r.get("rating", ""),
                r.get("reviews_count", ""), r.get("follow_up_count", 0), r.get("last_follow_up_at", ""), r.get("current_lead_status", "")
            ])

    csv_path = safe_write_text_or_csv(csv_path, write_standard_csv, mode="w", encoding="utf-8-sig", newline="")

    # 5. Professional High-Speed Excel Export
    final_xlsx_path = None
    if HAS_OPENPYXL:
        try:
            wb = openpyxl.Workbook()
            ws1 = wb.active
            if ws1 is None:
                ws1 = wb.create_sheet(title="Delivered Emails")
            else:
                ws1.title = "Delivered Emails"

            # Append headers
            ws1.append(headers)

            # Predefined optimal column widths for lightning speed
            col_widths = [10, 10, 28, 28, 26, 18, 16, 36, 16, 22, 18, 16, 20, 28, 18, 10, 14, 16, 22, 20]
            for col_idx, width in enumerate(col_widths, 1):
                ws1.column_dimensions[get_column_letter(col_idx)].width = width

            # Batch append rows
            for r in rows:
                ws1.append([
                    r.get("log_id", ""), r.get("lead_id", ""), r.get("business_name", ""), r.get("recipient_email_clean", ""), r.get("sender_email", ""),
                    r.get("campaign", ""), r.get("pitch_type", ""), r.get("subject", ""), r.get("delivery_status", ""), r.get("sent_at", ""),
                    r.get("city", ""), r.get("country", ""), r.get("niche", ""), r.get("website_url", ""), r.get("phone", ""), r.get("rating", ""),
                    r.get("reviews_count", ""), r.get("follow_up_count", 0), r.get("last_follow_up_at", ""), r.get("current_lead_status", "")
                ])

            # Apply professional Excel Table styling natively
            total_rows = max(len(rows) + 1, 2)
            tab_ref = f"A1:{get_column_letter(len(headers))}{total_rows}"
            table = Table(displayName="DeliveredEmailsTable", ref=tab_ref)
            style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
            table.tableStyleInfo = style
            ws1.add_table(table)
            ws1.freeze_panes = "A2"

            # Sheet 2: Campaign Summary
            ws2 = wb.create_sheet(title="Campaign Summary")
            ws2.cell(row=1, column=1, value="ALORIA HUNTER — OUTREACH CAMPAIGN SUMMARY").font = Font(name="Calibri", size=14, bold=True, color="0F172A")
            ws2.cell(row=2, column=1, value=f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}").font = Font(name="Calibri", size=10, italic=True, color="64748B")

            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            border_thin = Border(
                left=Side(style='thin', color='E2E8F0'),
                right=Side(style='thin', color='E2E8F0'),
                top=Side(style='thin', color='E2E8F0'),
                bottom=Side(style='thin', color='E2E8F0')
            )

            summary_data = [
                ("Total Delivered Email Events", len(rows)),
                ("Unique Recipient Businesses", len(sorted_unique_emails)),
                ("Initial Cold Pitches", sum(1 for r in rows if r.get("pitch_type") == "INITIAL")),
                ("Follow-up #1 Dispatched", sum(1 for r in rows if r.get("pitch_type") == "FOLLOW_UP_1")),
                ("Follow-up #2 Dispatched", sum(1 for r in rows if r.get("pitch_type") == "FOLLOW_UP_2")),
            ]
            # Dynamically include all campaigns present
            for c_name, count in sorted(campaign_counts.items(), key=lambda x: -x[1]):
                summary_data.append((f"Campaign: {c_name}", count))

            ws2.cell(row=4, column=1, value="Metric").font = header_font
            ws2.cell(row=4, column=1).fill = header_fill
            ws2.cell(row=4, column=2, value="Count").font = header_font
            ws2.cell(row=4, column=2).fill = header_fill

            for r_idx, (metric, val) in enumerate(summary_data, 5):
                c1 = ws2.cell(row=r_idx, column=1, value=metric)
                c2 = ws2.cell(row=r_idx, column=2, value=val)
                c1.border = border_thin
                c2.border = border_thin
                c2.alignment = Alignment(horizontal="right")
                if r_idx % 2 == 0:
                    c1.fill = alt_fill
                    c2.fill = alt_fill

            ws2.column_dimensions["A"].width = 38
            ws2.column_dimensions["B"].width = 18

            # Safe save with PermissionError fallback
            try:
                wb.save(xlsx_path)
                final_xlsx_path = xlsx_path
            except PermissionError:
                ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                alt_xlsx = out_dir / f"outreach_campaign_sent_emails_{ts}.xlsx"
                wb.save(alt_xlsx)
                print(f"[!] Warning: '{xlsx_path.name}' is currently locked. Saved to '{alt_xlsx.name}' instead.")
                final_xlsx_path = alt_xlsx
        except Exception as ex:
            print(f"[-] Warning: Failed generating Excel workbook: {ex}")
            final_xlsx_path = None
    else:
        print("[!] Note: openpyxl is not installed. Excel export skipped (CSV, JSON & TXT available).")

    return {
        "total_sent": len(rows),
        "unique_emails": len(sorted_unique_emails),
        "xlsx_path": str(final_xlsx_path) if final_xlsx_path else str(xlsx_path),
        "csv_path": str(csv_path),
        "agent_json_path": str(agent_json_path),
        "agent_csv_path": str(agent_csv_path),
        "agent_txt_path": str(agent_txt_path)
    }

def main():
    parser = argparse.ArgumentParser(description="Aloria Hunter - Sent Emails & Outreach Exporter")
    parser.add_argument("--output-dir", "-o", default=None, help="Directory to save export files")
    parser.add_argument("--business", "-b", default=None, help="Filter by campaign / business_id")
    parser.add_argument("--quiet", "-q", action="store_true", help="Suppress console output")
    args = parser.parse_args()

    res = generate_outreach_exports(output_dir=args.output_dir, business_id=args.business)
    if not args.quiet:
        print(f"\n[OK] Export complete: {res['total_sent']} emails ({res['unique_emails']} unique contacted leads)")
        print(f"  * JSON:  {res['agent_json_path']}")
        print(f"  * CSV:   {res['csv_path']}")
        print(f"  * Agent: {res['agent_csv_path']}")
        print(f"  * TXT:   {res['agent_txt_path']}")
        if res.get("xlsx_path"):
            print(f"  * Excel: {res['xlsx_path']}")
        print()

if __name__ == "__main__":
    main()
